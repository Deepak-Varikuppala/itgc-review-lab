"""Builds the logical access testing workpaper in Excel."""

from __future__ import annotations

import hashlib
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from findings import FINDINGS, OBSERVATIONS

NAVY = "1F3864"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=15, color=NAVY)
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
FILLS = {
    "Pass": PatternFill("solid", fgColor="C6EFCE"),
    "Exception": PatternFill("solid", fgColor="FFC7CE"),
    "Not applicable": PatternFill("solid", fgColor="E7E6E6"),
    "High": PatternFill("solid", fgColor="FFC7CE"),
    "Medium": PatternFill("solid", fgColor="FFEB9C"),
    "Low": PatternFill("solid", fgColor="DDEBF7"),
}


def _print_setup(ws, header_row=None):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    if header_row:
        ws.print_title_rows = f"{header_row}:{header_row}"
    ws.oddFooter.center.text = "&A  |  Page &P of &N"


def _table(ws, start_row, headers, rows, widths, fill_cols=()):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, row in enumerate(rows, start_row + 1):
        for c, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.alignment, cell.border = WRAP, BORDER
            if c in fill_cols and value in FILLS:
                cell.fill = FILLS[value]
            elif c in fill_cols and isinstance(value, str):
                if value.startswith("Effective"):
                    cell.fill = FILLS["Pass"]
                elif value.startswith(("Not operating", "Design deficiency")):
                    cell.fill = FILLS["Exception"]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if rows:
        ws.auto_filter.ref = (
            f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"
        )
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    _print_setup(ws, start_row)


def _label_block(ws, rows, start=1):
    for i, (label, value) in enumerate(rows, start):
        ws.cell(row=i, column=1, value=label).font = BOLD
        ws.cell(row=i, column=1).alignment = WRAP
        cell = ws.cell(row=i, column=2, value=value)
        cell.alignment = WRAP
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
        lines = max(1, len(str(value)) // 95 + 1)
        ws.row_dimensions[i].height = 15 * lines
    return start + len(rows)


def build_workpaper(ev, tests, out: Path) -> Path:
    wb = Workbook()
    wb.properties.creator = "Deepak Kumar Varikuppala"
    wb.properties.title = "Northwind Retail: Logical Access ITGC Review"

    ws = wb.active
    ws.title = "Cover"
    ws["A1"] = "Northwind Retail: Logical Access ITGC Review"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")
    files = sorted(p for p in ev.dir.iterdir() if p.is_file())
    _label_block(
        ws,
        [
            ("Entity", "Northwind Retail (fictional company, lab environment)"),
            ("System", "Active Directory domain northwind.internal on linux.northwind.internal"),
            ("Audit period", "1 July 2026 to 30 September 2026"),
            ("Evidence extracted", f"{ev.as_at} (folder {ev.dir.name})"),
            ("Criteria", "Northwind Retail Access Control Policy v2.1, approved 15 January 2026"),
            (
                "Approach",
                "Each test covers the whole population, so no sampling was used. Directory "
                "evidence was extracted read only by the system administrator, checked "
                "for completeness (record counts) and integrity (SHA-256), and compared "
                "with HR, contractor, access request and exception records.",
            ),
            ("Tests performed", f"{len(tests)}"),
            ("Tests with exceptions", f"{sum(1 for t in tests if t.exceptions())}"),
            (
                "Findings",
                f"{len(FINDINGS)} ("
                + ", ".join(
                    f"{sum(1 for f in FINDINGS if f.rating == r)} {r}"
                    for r in ("High", "Medium", "Low")
                )
                + ")",
            ),
            ("Prepared by", "Deepak Kumar Varikuppala"),
            ("Workpaper generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ],
        start=3,
    )
    ws.column_dimensions["A"].width = 24
    for col in "BCD":
        ws.column_dimensions[col].width = 40
    _print_setup(ws)

    ws = wb.create_sheet("Summary")
    rows = []
    for t in tests:
        tested = [i for i in t.items if i.result != "Not applicable"]
        rows.append(
            [
                t.test_id,
                t.title,
                t.policy_ref,
                t.test_type,
                len(tested),
                len(t.exceptions()),
                t.conclusion,
                ", ".join(t.findings),
                t.iso27001,
                t.cis_controls,
            ]
        )
    _table(
        ws,
        1,
        [
            "Test",
            "Control tested",
            "Policy",
            "Type",
            "Items tested",
            "Exceptions",
            "Conclusion",
            "Findings",
            "ISO/IEC 27001:2022",
            "CIS Controls v8.1",
        ],
        rows,
        [8, 46, 10, 18, 9, 10, 36, 9, 30, 32],
        fill_cols=(7,),
    )

    ws = wb.create_sheet("Findings")
    rows = [
        [
            f.finding_id,
            f.title,
            f.rating,
            ", ".join(f.tests),
            f.policy,
            f.condition,
            f.recommendation,
            "",
            "",
            "",
        ]
        for f in FINDINGS
    ]
    rows += [[o[0], o[1], "Observation", "", "", o[2], "", "", "", ""] for o in OBSERVATIONS]
    _table(
        ws,
        1,
        [
            "ID",
            "Finding",
            "Rating",
            "Tests",
            "Policy",
            "Condition",
            "Recommendation",
            "Owner",
            "Management response",
            "Target date",
        ],
        rows,
        [6, 34, 11, 12, 9, 60, 50, 14, 30, 12],
        fill_cols=(3,),
    )

    for t in tests:
        ws = wb.create_sheet(t.test_id)
        ws["A1"] = f"{t.test_id}: {t.title}"
        ws["A1"].font = TITLE_FONT
        ws.merge_cells("A1:D1")
        nxt = _label_block(
            ws,
            [
                ("Policy clause", t.policy_ref),
                ("ISO/IEC 27001:2022", t.iso27001),
                ("CIS Controls v8.1", t.cis_controls),
                ("Test type", t.test_type),
                ("Control objective", t.objective),
                ("Test procedure", t.procedure),
                ("Population", t.population),
                ("Sample", "Whole population"),
                ("Conclusion", t.conclusion),
                ("Findings raised", ", ".join(t.findings) or "None"),
            ],
            start=3,
        )
        conclusion_row = nxt - 2
        ws.cell(row=conclusion_row, column=2).fill = (
            FILLS["Pass"] if t.conclusion.startswith("Effective") else FILLS["Exception"]
        )
        _table(
            ws,
            nxt + 1,
            ["Item", "Detail", "Result", "Comment"],
            [[i.item, i.detail, i.result, i.comment] for i in t.items],
            [34, 34, 13, 70],
            fill_cols=(3,),
        )

    ws = wb.create_sheet("Evidence")
    rows = [[p.name, p.stat().st_size, hashlib.sha256(p.read_bytes()).hexdigest()] for p in files]
    _table(ws, 1, ["Evidence file", "Bytes", "SHA-256"], rows, [26, 10, 70])

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def build_change_ops_workpaper(ev, tests, samples, out: Path, seed: int) -> Path:
    """Workpaper for the change management, IT operations and server access review."""
    from change_ops_findings import FINDINGS as CO_FINDINGS
    from change_ops_findings import OBSERVATIONS as CO_OBSERVATIONS

    wb = Workbook()
    wb.properties.creator = "Deepak Kumar Varikuppala"
    wb.properties.title = "Northwind Retail: Change Management, IT Operations and Server Access"

    ws = wb.active
    ws.title = "Cover"
    ws["A1"] = "Northwind Retail: Change Management, IT Operations and Server Access Review"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")
    folders = [ev.app_dir, *(h.dir for h in ev.hosts.values()), ev.directory_dir]
    if ev.pc:
        folders.append(ev.pc.dir)
    _label_block(
        ws,
        [
            ("Entity", "Northwind Retail (fictional company, lab environment)"),
            (
                "Systems",
                "Finance application and its code repository (Gitea) on linux2; Linux "
                "servers linux (domain controller) and linux2 (application server)",
            ),
            ("Audit period", "1 July 2026 to 30 September 2026"),
            ("Evidence extracted", f"{ev.as_at} (" + ", ".join(f.name for f in folders) + ")"),
            (
                "Criteria",
                "Northwind Retail Change Management and IT Operations Standard v1.4 and "
                "Access Control Policy v2.1, both approved 15 January 2026",
            ),
            (
                "Approach",
                "Tests that compare system records without judgement cover the whole "
                "population. CM-04 and AC-18 need each item to be read and judged, so a "
                f"random sample was taken with Python's random.Random({seed}).sample(), with "
                "the size set by the table in the test plan. See the Samples sheet.",
            ),
            ("Tests performed", f"{len(tests)}"),
            ("Tests with exceptions", f"{sum(1 for t in tests if t.exceptions())}"),
            (
                "Findings",
                f"{len(CO_FINDINGS)} ("
                + ", ".join(
                    f"{sum(1 for f in CO_FINDINGS if f.rating == r)} {r}"
                    for r in ("High", "Medium", "Low")
                )
                + ")",
            ),
            ("Prepared by", "Deepak Kumar Varikuppala"),
            ("Workpaper generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ],
        start=3,
    )
    ws.column_dimensions["A"].width = 24
    for col in "BCD":
        ws.column_dimensions[col].width = 40
    _print_setup(ws)

    ws = wb.create_sheet("Summary")
    rows = [
        [
            t.test_id,
            t.title,
            t.policy_ref,
            t.test_type,
            t.sample or "Whole population",
            len([i for i in t.items if i.result != "Not applicable"]),
            len(t.exceptions()),
            t.conclusion,
            ", ".join(t.findings),
            t.iso27001,
            t.cis_controls,
        ]
        for t in tests
    ]
    _table(
        ws,
        1,
        [
            "Test",
            "Control tested",
            "Criteria",
            "Type",
            "Sample",
            "Items tested",
            "Exceptions",
            "Conclusion",
            "Findings",
            "ISO/IEC 27001:2022",
            "CIS Controls v8.1",
        ],
        rows,
        [8, 44, 16, 18, 22, 9, 10, 34, 14, 30, 32],
        fill_cols=(8,),
    )

    ws = wb.create_sheet("Findings")
    rows = [
        [
            f.finding_id,
            f.title,
            f.rating,
            ", ".join(f.tests),
            f.policy,
            f.condition,
            f.recommendation,
            "",
            "",
            "",
        ]
        for f in CO_FINDINGS
    ]
    rows += [[o[0], o[1], "Observation", "", "", o[2], "", "", "", ""] for o in CO_OBSERVATIONS]
    _table(
        ws,
        1,
        [
            "ID",
            "Finding",
            "Rating",
            "Tests",
            "Criteria",
            "Condition",
            "Recommendation",
            "Owner",
            "Management response",
            "Target date",
        ],
        rows,
        [6, 34, 11, 16, 16, 60, 50, 14, 30, 12],
        fill_cols=(3,),
    )

    ws = wb.create_sheet("Samples")
    rows = []
    for test_id, s in samples.items():
        chosen = set(s["sample"])
        for n, key in enumerate(s["population"], 1):
            rows.append([test_id, n, key, "Selected" if key in chosen else ""])
    _table(
        ws,
        1,
        ["Test", "Population item", "Item", f"Selected (seed {seed})"],
        rows,
        [8, 10, 46, 18],
    )

    for t in tests:
        ws = wb.create_sheet(t.test_id)
        ws["A1"] = f"{t.test_id}: {t.title}"
        ws["A1"].font = TITLE_FONT
        ws.merge_cells("A1:D1")
        nxt = _label_block(
            ws,
            [
                ("Criteria", t.policy_ref),
                ("ISO/IEC 27001:2022", t.iso27001),
                ("CIS Controls v8.1", t.cis_controls),
                ("Test type", t.test_type),
                ("Control objective", t.objective),
                ("Test procedure", t.procedure),
                ("Population", t.population),
                ("Sample", t.sample or "Whole population"),
                ("Conclusion", t.conclusion),
                ("Findings raised", ", ".join(t.findings) or "None"),
            ],
            start=3,
        )
        ws.cell(row=nxt - 2, column=2).fill = (
            FILLS["Pass"] if t.conclusion.startswith("Effective") else FILLS["Exception"]
        )
        _table(
            ws,
            nxt + 1,
            ["Item", "Detail", "Result", "Comment"],
            [[i.item, i.detail, i.result, i.comment] for i in t.items],
            [34, 40, 13, 70],
            fill_cols=(3,),
        )

    ws = wb.create_sheet("Evidence")
    rows = []
    for folder in folders:
        for p in sorted(x for x in folder.iterdir() if x.is_file()):
            rows.append(
                [folder.name, p.name, p.stat().st_size, hashlib.sha256(p.read_bytes()).hexdigest()]
            )
    _table(ws, 1, ["Folder", "Evidence file", "Bytes", "SHA-256"], rows, [32, 30, 10, 70])

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
